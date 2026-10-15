"""Compare the old renderer and new renderer on identical, illustrative slides.

Usage: python examples/benchmark_design_engine.py /path/to/previous-project.zip
Scores here are objective structural metrics, NOT human-rated visual quality.
"""
from __future__ import annotations
from pathlib import Path
from io import BytesIO
import sys, zipfile, types, json, argparse
from pptx import Presentation
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
from app.engine import render as new_render
from app.quality import audit_pptx
from make_quality_demo import SLIDES

def metrics(data):
    deck=Presentation(BytesIO(data))
    shape_sequences=[]
    native=0;connectors=0;text_boxes=0
    for slide in deck.slides:
        ids=[]
        for shape in slide.shapes:
            if shape.has_chart:native+=1
            if shape.shape_type==9:connectors+=1
            if shape.has_text_frame and shape.text.strip():text_boxes+=1
            ids.append(str(shape.shape_type))
        shape_sequences.append('|'.join(ids))
    audit=audit_pptx(data,verify_external=True)
    return {'slides':len(deck.slides),'native_charts':native,'native_connectors':connectors,
            'text_elements':text_boxes,'distinct_shape_sequences':len(set(shape_sequences)),
            'critical_geometry_findings':audit['critical_count'],
            'heuristic_text_and_readability_warnings':audit['warning_count'],
            'external_pdf_reopen':audit['independent_render']['status']}

def main(old_archive):
    with zipfile.ZipFile(old_archive) as z:
        code=z.read('pptx-ai/backend/app/engine.py').decode()
    old=types.ModuleType('old_renderer_fixture');exec(compile(code,'old_engine.py','exec'),old.__dict__)
    a=old.render('Design benchmark',SLIDES,'midnight');b=new_render('Design benchmark',SLIDES,'midnight')
    report={'fixture':'15 illustrative slides with user-provided dummy chart data; not a model quality study',
            'baseline_source':str(Path(old_archive).name),
            'measured':'PPTX structure and independent LibreOffice PDF rendering only',
            'old':metrics(a),'new':metrics(b),
            'human_visual_quality':'not_measured', 'model_factual_accuracy':'not_measured',
            'generation_cost':'not_measured','render_speed':'not_measured'}
    Path(__file__).with_name('benchmark_report.json').write_text(json.dumps(report,indent=2))
    print(json.dumps(report,indent=2))

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('previous_project_zip');args=p.parse_args()
    main(args.previous_project_zip)
