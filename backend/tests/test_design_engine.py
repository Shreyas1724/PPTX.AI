"""Regression suite for native editable layouts, validation and quality gate."""
from io import BytesIO
import pytest
from pydantic import ValidationError
from pptx import Presentation
from pptx.util import Inches
from app.engine import render,local_slides
from app.slide_spec import SlideSpec,ChartSpec
from app.design_planner import plan_layouts
from app.quality import audit_pptx
from test_identity import setup,signup


def test_legacy_shape_backward_compatible():
    slides=local_slides('Low-cost sensors',3)
    output=render('Sensors',slides,'light')
    deck=Presentation(BytesIO(output))
    assert len(deck.slides)==3
    assert any('Overview' in sh.text for sh in deck.slides[1].shapes if sh.has_text_frame)


def test_planner_matches_content_and_preserves_override():
    src=[SlideSpec(title='My introduction',body=['Context']),
         SlideSpec(title='Technical Architecture',body=['Browser','API','Database']),
         SlideSpec(title='Measured data',chart={'kind':'bar','points':[{'label':'A','value':1},{'label':'B','value':3}]}),
         SlideSpec(title='Results',layout='quote',body=['Human-approved statement'])]
    chosen=plan_layouts(src)
    assert [slide.layout for slide in chosen]==['title','architecture','chart','quote']


def test_invalid_chart_is_rejected_without_fallback_to_fake_content():
    with pytest.raises(ValidationError):
        SlideSpec(title='Incorrect data',chart={'kind':'pie','points':[{'label':'A','value':-5},{'label':'B','value':5}]})
    with pytest.raises(ValidationError):
        SlideSpec(title='Invalid layout',layout='magic-animation')
    with pytest.raises(ValidationError):
        SlideSpec(title='Unsafe',body=['x'*300])


def test_native_charts_and_editable_text():
    kinds=['bar','line','pie']
    slides=[{'title':f'{k} example','layout':'chart','chart':{'kind':k,
             'points':[{'label':'A','value':1},{'label':'B','value':3},{'label':'C','value':2}]}} for k in kinds]
    output=render('Test native charts',slides,'midnight')
    deck=Presentation(BytesIO(output))
    assert len(deck.slides)==3
    assert all(sum(shape.has_chart for shape in slide.shapes)==1 for slide in deck.slides)
    assert sum(shape.has_text_frame for slide in deck.slides for shape in slide.shapes)>6
    for slide in deck.slides:
        chart=next(x.chart for x in slide.shapes if x.has_chart)
        assert chart.plots[0].series[0].values==(1.0,3.0,2.0)


def test_diagrams_have_editable_elements():
    content=[{'title':'Architecture','layout':'architecture','steps':['Client','Service','Database']},
             {'title':'Workflow','layout':'process','steps':['Input','Plan','Build','Review','Export']},
             {'title':'Delivery timeline','layout':'timeline','steps':['Design','Implement','Test']}]
    deck=Presentation(BytesIO(render('Demo',content,'aurora')))
    assert len(deck.slides)==3
    assert all(sum(sh.has_text_frame for sh in slide.shapes)>6 for slide in deck.slides)
    assert any(sh.shape_type==9 for slide in deck.slides for sh in slide.shapes)  # connector


def test_all_layout_identifiers_render_without_outside_canvas():
    from typing import get_args
    from app.slide_spec import Layout
    items=[]
    for i,layout in enumerate(get_args(Layout)):
        item={'title':f'Example for {layout}','layout':layout,'body':['First important point','Second explanation']}
        if layout=='chart': item['chart']={'kind':'bar','points':[{'label':'One','value':3},{'label':'Two','value':7}]}
        if layout=='kpi': item['metrics']=[{'label':'Measured uptime','value':'99.0%'}]
        items.append(item)
    # API deliberately caps a deck at 20, so run across batches.
    for start in range(0,len(items),12):
        result=audit_pptx(render('Layout fixture',items[start:start+12],'light'))
        assert result['critical_count']==0,result['findings']


def test_quality_detects_bad_geometry():
    from pptx import Presentation as P
    from pptx.util import Inches
    p=P();slide=p.slides.add_slide(p.slide_layouts[6]);slide.shapes.add_textbox(Inches(14),Inches(1),Inches(2),Inches(1)).text='Off screen'
    buf=BytesIO();p.save(buf)
    result=audit_pptx(buf.getvalue())
    assert result['status']=='failed'
    assert any(f['code']=='outside_canvas' for f in result['findings'])


def test_editor_quality_endpoint_is_private(setup):
    client,codes=setup
    assert client.get('/api/v1/presentations/fake/quality').status_code in (401,403)
    signup(client,codes)
    project=client.post('/api/v1/projects',json={'name':'Report test'}).json()['id']
    deck=client.post('/api/v1/presentations',json={'project_id':project,'title':'Geometry report'}).json()['id']
    sample={'title':'Native editable chart','body':['Read this sample carefully.'], 'layout':'chart',
            'chart':{'kind':'bar','points':[{'label':'2024','value':5},{'label':'2025','value':8}]}}
    assert client.post('/api/v1/presentations/'+deck+'/slides',json=sample).status_code==200
    quality=client.get('/api/v1/presentations/'+deck+'/quality');assert quality.status_code==200,quality.text
    assert quality.json()['slide_count']==1 and quality.json()['critical_count']==0
    assert client.get('/api/v1/presentations/invalid/quality').status_code==404
    client.post('/api/v1/auth/logout');signup(client,codes,2)
    assert client.get('/api/v1/presentations/'+deck+'/quality').status_code==404


def test_visual_chart_context_preserved_in_notes():
    slides=[{'title':'Example','layout':'chart','body':['First main conclusion','Secondary supporting claim'],
             'chart':{'kind':'bar','points':[{'label':'A','value':1},{'label':'B','value':2}]}}]
    deck=Presentation(BytesIO(render('Chart notes',slides,'midnight')))
    note_text=deck.slides[0].notes_slide.notes_text_frame.text
    assert 'Secondary supporting claim' in note_text
