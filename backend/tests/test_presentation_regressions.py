"""Regressions for exported content, provider failures and document parsing."""
from io import BytesIO

import httpx
import pytest
from pptx import Presentation
from pptx.util import Inches, Pt

from app import intelligence, quality
from app.documents import parse
from app.engine import render
from app.design_planner import plan_layouts
from app.slide_spec import SlideSpec


@pytest.mark.parametrize(('title', 'layout'), [
    ('References', 'references'), ('Bibliography', 'references'),
    ('Comparison of options', 'comparison'), ('Compare approaches', 'comparison'),
    ('A vs. B', 'comparison'),
])
def test_planner_recognizes_complete_keywords(title, layout):
    slides = [SlideSpec(title='Introduction'), SlideSpec(title=title, body=['First', 'Second'])]
    assert plan_layouts(slides)[1].layout == layout


def test_planner_recognizes_plural_timeline_steps():
    slide = SlideSpec(title='Project milestones', steps=['Design', 'Build'])
    assert plan_layouts([slide])[0].layout == 'timeline'


def test_document_parser_is_independent_of_working_directory(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv('PYTHONPATH', raising=False)
    assert parse(b'Reference text from another directory', 'txt') == 'Reference text from another directory'


def test_malformed_provider_response_retries_before_success(monkeypatch):
    monkeypatch.setenv('LLM_API_KEY', 'test-provider-key')
    responses = iter([
        {'choices': []},
        {'choices': [{'message': {'content': '{"slides":[{"title":"Recovered","body":["Verified point"]}]}'}}]},
    ])
    calls = []

    class Client:
        def __init__(self, **kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

        def post(self, url, **kwargs):
            calls.append(url)
            return httpx.Response(200, request=httpx.Request('POST', url), json=next(responses))

    monkeypatch.setattr(intelligence.httpx, 'Client', Client)
    deck, usage = intelligence.llm('Validate this deck', {}, intelligence.DeckIR, 1)
    assert deck.slides[0].title == 'Recovered'
    assert len(calls) == 2 and usage['attempts'] == 2


def test_quality_respects_paragraph_font_inheritance():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    small = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    small.text = 'Small inherited font'
    small.text_frame.paragraphs[0].font.size = Pt(9)
    large = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(1), Inches(.4))
    large.text = 'A very large heading that cannot fit'
    large.text_frame.paragraphs[0].font.size = Pt(60)
    output = BytesIO()
    deck.save(output)
    codes = {finding['code'] for finding in quality.audit_pptx(output.getvalue())['findings']}
    assert {'small_font', 'possible_text_overflow'} <= codes


def test_process_cards_and_chart_units_pass_readability_checks():
    slides = [
        {'title': 'Methodology', 'layout': 'process', 'steps': [
            'Collect user materials', 'Separate facts from assumptions', 'Draft an outline', 'Validate key claims',
        ]},
        {'title': 'Measured data', 'layout': 'chart', 'chart': {
            'kind': 'bar', 'unit': 'sample units',
            'points': [{'label': 'A', 'value': 1}, {'label': 'B', 'value': 2}],
        }},
    ]
    result = quality.audit_pptx(render('Readable data', slides, 'light'))
    assert result['warning_count'] == 0, result['findings']


def test_failed_independent_render_fails_overall_quality(monkeypatch):
    monkeypatch.setattr(quality, 'verify_render', lambda *_: {'status': 'failed', 'pages': None, 'detail': 'Conversion failed'})
    result = quality.audit_pptx(render('Example', [{'title': 'Example'}], 'light'), verify_external=True)
    assert result['critical_count'] == 0
    assert result['status'] == 'failed'


def test_unreadable_independent_pdf_is_reported_as_failure(monkeypatch):
    from pathlib import Path
    from types import SimpleNamespace

    monkeypatch.setattr(quality.shutil, 'which', lambda _: '/usr/bin/libreoffice')

    def convert(command, **kwargs):
        directory = Path(command[command.index('--outdir') + 1])
        (directory / 'input.pdf').write_bytes(b'')
        return SimpleNamespace(returncode=0)

    monkeypatch.setattr(quality.subprocess, 'run', convert)
    result = quality.verify_render(b'example', 1)
    assert result['status'] == 'failed'
    assert result['detail'] == 'EmptyFileError'


@pytest.mark.parametrize('layout', ['chart', 'kpi'])
def test_visual_layout_has_one_insight_and_preserves_all_sources(layout):
    slide = {
        'title': 'Measured results', 'layout': layout,
        'key_message': 'The central measured result',
        'body': ['Supporting detail one', 'Supporting detail two'],
        'sources': ['Source one: ' + 'x' * 150, 'Source two'],
    }
    if layout == 'chart':
        slide['chart'] = {'kind': 'bar', 'points': [{'label': 'A', 'value': 1}, {'label': 'B', 'value': 2}]}
    else:
        slide['metrics'] = [{'label': 'Measured total', 'value': '2'}]
    exported = Presentation(BytesIO(render('Results', [slide], 'midnight'))).slides[0]
    insights = [shape for shape in exported.shapes if shape.has_text_frame and 6.3 <= shape.top / 914400 <= 6.8]
    assert len(insights) == 1
    assert slide['key_message'] in insights[0].text
    notes = exported.notes_slide.notes_text_frame.text
    assert all(text in notes for text in slide['body'] + slide['sources'])


def test_missing_job_refunds_its_reservation_once(monkeypatch):
    from sqlalchemy import create_engine, func, select
    from sqlalchemy.orm import sessionmaker
    from app import core, worker

    engine = create_engine('sqlite://')
    core.Base.metadata.create_all(engine)
    sessions = sessionmaker(bind=engine, expire_on_commit=False)
    monkeypatch.setattr(worker, 'Session', sessions)
    try:
        with sessions() as session:
            user = core.User(email='orphan@example.com', password='unused', credits=28)
            session.add(user)
            session.flush()
            operation = core.Operation(user_id=user.id, job_id='removed-job', request_key='orphan', fingerprint='test',
                                       kind='generation', payload='{}', cost=2, base_version=0)
            session.add(operation)
            session.commit()
            user_id, operation_id = user.id, operation.id
        worker.process('removed-job')
        worker.process('removed-job')
        with sessions() as session:
            assert session.get(core.User, user_id).credits == 30
            assert session.get(core.Operation, operation_id).state == 'refunded'
            assert session.scalar(select(func.count()).select_from(core.CreditEvent)) == 1
    finally:
        engine.dispose()
