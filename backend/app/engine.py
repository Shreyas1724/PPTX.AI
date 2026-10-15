"""PPTX.AI native PowerPoint design engine.

Backward-compatible entry points: local_slides(prompt,count), render(title,slides,theme).
Every visible element (except optional external images, not implemented here) is
an editable PowerPoint shape, text box, connector or native chart.
"""
from __future__ import annotations
from io import BytesIO
import math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_TICK_MARK
from .design_system import get_theme
from .slide_spec import validate_deck, SlideSpec
from .design_planner import plan_layouts

W,H=13.333,7.5

def local_slides(prompt,count,context=''):
    """A clearly labeled locally produced outline, never simulated research."""
    title=prompt.strip()[:75] or 'Untitled presentation'
    headings=['Overview','Background','Key ideas','How it works','Opportunities','Challenges','Approach','Examples','Next steps','Conclusion']
    return [{'title':title if i==0 else headings[(i-1)%len(headings)],
             'body':([f'Introduction to {title}','Purpose, audience, and scope'] if i==0 else
                     [f'Explore {title.lower()} through {headings[(i-1)%len(headings)].lower()}.',
                      'Add verified facts and examples before publishing.']),
             'notes':'Unresearched draft outline. Review all facts before sharing.'}
            for i in range(count)]

def rect(sl,x,y,w,h,th,fill='panel',radius=False,line=None):
    sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
            Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid();sh.fill.fore_color.rgb=th.color(fill)
    if line:sh.line.color.rgb=th.color(line);sh.line.width=Pt(0.7)
    else:sh.line.fill.background()
    if radius:
        try:sh.adjustments[0]=.13
        except (IndexError,AttributeError):pass
    return sh

def label(sl,text,x,y,w,h,th,pt=18,weight=False,color='ink',align=None,
          vertical=MSO_ANCHOR.MIDDLE,margin=0.02):
    sh=sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame;tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin);tf.margin_top=tf.margin_bottom=Inches(.02)
    tf.vertical_anchor=vertical
    p=tf.paragraphs[0];p.text=str(text);p.font.name=th.heading_font if weight else th.body_font
    p.font.size=Pt(pt);p.font.bold=weight;p.font.color.rgb=th.color(color)
    if align is not None:p.alignment=align
    return sh

def wrapped(sl,body,x,y,w,h,th,pt=19,numbered=False,color='ink'):
    """Separate editable rows give predictable placement and allow QA to measure each item."""
    if not body:return
    n=len(body)
    gap=min(.18,.55/max(n,1));usable=h-gap*(n-1);rh=usable/n
    for j,line_text in enumerate(body):
        size=min(pt,18 if len(line_text)>130 else pt)
        if rh < .7:size=min(size,17)
        prefix=f'{j+1:02d}' if numbered else '•'
        label(sl,prefix,x,y+j*(rh+gap),.48,min(rh,.43),th,pt=12,color='accent',weight=True)
        label(sl,line_text,x+.52,y+j*(rh+gap),w-.56,rh,th,pt=size,vertical=MSO_ANCHOR.MIDDLE)

def base(deck,i,n,slide,theme):
    th=get_theme(theme)
    sl=deck.slides.add_slide(deck.slide_layouts[6])
    sl.background.fill.solid();sl.background.fill.fore_color.rgb=th.color('background')
    rect(sl,0,0,.10,H,th,'accent')
    label(sl,'PPTX.AI  /  STUDIO',.65,.25,3,.28,th,10,True,'accent')
    label(sl,f'{i+1:02d}  /  {n:02d}',11.67,7.02,1,.20,th,10,False,'secondary',align=PP_ALIGN.RIGHT)
    rect(sl,.65,6.88,12,.012,th,'divider')
    if slide.sources:
        label(sl,'SOURCE  '+str(slide.sources[0])[:105],.68,6.96,10.7,.3,th,9,color='secondary')
    return sl,th

def heading(sl,slide,th,title_size=31):
    label(sl,slide.title,.78,.73,11.83,.85,th,pt=title_size,weight=True)
    if slide.subtitle:label(sl,slide.subtitle,.82,1.58,11.3,.42,th,pt=15,color='secondary')
    if slide.key_message and not (slide.layout=='chart' and slide.chart or slide.layout=='kpi' and slide.metrics):
        label(sl,slide.key_message,.82,6.40,11.4,.35,th,pt=12,color='accent')

def card(sl,th,x,y,w,h,index,title,desc='',tint=False):
    rect(sl,x,y,w,h,th,'panel_alt' if tint else 'panel',radius=True)
    label(sl,f'{index:02d}' if isinstance(index,int) else index,x+.24,y+.23,.55,.32,th,13,True,'accent')
    title_top=.67 if h>1.4 else .60
    title_height=min(.85,h*.34) if desc else max(.3,h-title_top-.16)
    title_size=(21 if len(title)<66 else 16) if h>1.65 else (17 if h>1.4 else 14)
    label(sl,title,x+.25,y+title_top,w-.5,title_height,th,title_size,True)
    if desc and h>1.65:label(sl,desc,x+.25,y+1.55,w-.5,h-1.70,th,14,color='secondary',vertical=MSO_ANCHOR.TOP)

def native_chart(sl,spec,th):
    chart=spec.chart
    data=CategoryChartData();data.categories=[p.label for p in chart.points]
    data.add_series(chart.title or 'Values',[p.value for p in chart.points])
    kind={'bar':XL_CHART_TYPE.COLUMN_CLUSTERED,'line':XL_CHART_TYPE.LINE_MARKERS,
          'pie':XL_CHART_TYPE.PIE}[chart.kind]
    chart_obj=sl.shapes.add_chart(kind,Inches(.9),Inches(2.03),Inches(11.48),Inches(4.42),data).chart
    chart_obj.has_legend=False
    chart_obj.has_title=False
    chart_obj.chart_style=10
    if chart.kind!='pie':
        chart_obj.value_axis.has_major_gridlines=True
        chart_obj.value_axis.tick_labels.font.size=Pt(12)
        chart_obj.value_axis.tick_labels.font.color.rgb=th.color('ink')
        chart_obj.category_axis.tick_labels.font.size=Pt(12)
        chart_obj.category_axis.tick_labels.font.color.rgb=th.color('ink')
        chart_obj.category_axis.major_tick_mark=XL_TICK_MARK.NONE
        chart_obj.value_axis.major_tick_mark=XL_TICK_MARK.NONE
    try:
        plot=chart_obj.plots[0];plot.has_data_labels=(chart.kind=='pie')
        if plot.has_data_labels:
            plot.data_labels.font.size=Pt(12)
            plot.data_labels.font.color.rgb=th.color('ink')
        if chart.kind=='bar':plot.series[0].format.fill.solid();plot.series[0].format.fill.fore_color.rgb=th.color('accent')
        if chart.kind=='line':plot.series[0].format.line.color.rgb=th.color('accent')
    except (AttributeError,IndexError):pass
    if chart.unit:label(sl,f'UNIT  {chart.unit}',10,1.73,2.25,.3,th,12,color='secondary',align=PP_ALIGN.RIGHT)

def connect(sl,x1,y1,x2,y2,th):
    line=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    line.line.color.rgb=th.color('accent_two');line.line.width=Pt(2)
    return line

def diagram(sl,items,th,style='process'):
    items=list(items)[:8];n=len(items)
    if n==0:return
    if style=='architecture':
        # Top-to-bottom editable hierarchy, with connectors between tiers.
        for idx,item in enumerate(items):
            y=1.90+idx*(4.58/n)
            rect(sl,2.07,y,9.2,max(.42,4.2/n),th,'panel_alt' if idx%2 else 'panel',True)
            label(sl,item,2.42,y,8.5,max(.36,4.2/n),th,pt=17 if n<=5 else 13,weight=idx==0)
            if idx:connect(sl,6.67,y-.16,6.67,y,th)
        return
    if style=='timeline':
        n=min(n,6);items=items[:n]
        for idx,item in enumerate(items):
            x=.89+idx*(11.48/n);cw=(11.48/n)-.13
            label(sl,f'{idx+1:02d}',x,2.00,cw,.56,th,25,True,'accent')
            rect(sl,x,2.86,cw,.07,th,'accent_two' if idx%2 else 'accent')
            rect(sl,x,3.13,cw,2.28,th,'panel',True)
            label(sl,item,x+.19,3.37,cw-.39,1.60,th,pt=17 if n<=4 else 14,weight=True,vertical=MSO_ANCHOR.TOP)
        return
    # Process with native connectors, avoiding connectors through text.
    per_row=3 if n>4 else min(n,4)
    for idx,item in enumerate(items):
        row=idx//per_row;col=idx%per_row
        rows=math.ceil(n/per_row);cw=11.45/per_row-.24
        x=.9+col*(11.45/per_row);y=2.00+row*(4.38/rows)
        h=min(1.62,3.82/rows)
        card(sl,th,x,y,cw,h,idx+1,item,tint=idx%2==1)
        if col>0:connect(sl,x-.20,y+h*.52,x-.03,y+h*.52,th)

def render_spec(sl,spec,i,n,th):
    mode=spec.layout or 'bullets';body=spec.body
    if mode=='title':
        rect(sl,.72,1.85,7.5,4.6,th,'panel',True)
        label(sl,'IDEAS, BEAUTIFULLY PRESENTED',1.06,2.19,6.75,.42,th,13,True,'accent')
        label(sl,spec.title,1.04,2.74,6.83,1.55,th,34 if len(spec.title)<55 else 27,True)
        label(sl,spec.subtitle or spec.key_message or (body[0] if body else 'A presentation prepared with PPTX.AI'),1.06,4.73,6.62,.98,th,17,color='secondary')
        rect(sl,8.65,1.85,3.80,4.6,th,'panel_alt',True)
        if body:
            wrapped(sl,body,9.00,2.26,3.05,3.85,th,pt=14,numbered=True)
        else:
            for j in range(4):rect(sl,9.05+j*.17,2.28+j*.5,2.7,1.08,th,'accent' if j==0 else 'panel',True)
        return
    if mode in ('section','conclusion','quote'):
        label(sl,'THE BIG PICTURE' if mode=='section' else 'KEY TAKEAWAY' if mode=='conclusion' else 'PERSPECTIVE',.95,1.88,7,.45,th,12,True,'accent')
        label(sl,spec.title,1.03,2.49,10.9,1.62,th,34,True)
        if body:wrapped(sl,body,1.08,4.27,10.5,2.22,th,pt=20)
        return
    heading(sl,spec,th)
    if mode=='chart' and spec.chart:
        native_chart(sl,spec,th)
        insight=spec.key_message or (body[0] if body else '')
        if insight:label(sl,'INSIGHT  '+insight,.95,6.42,11.25,.34,th,12,color='accent')
        return
    if mode=='kpi' and spec.metrics:
        num=len(spec.metrics);cw=11.68/num-.20
        for j,m in enumerate(spec.metrics):
            x=.82+j*(11.68/num)
            rect(sl,x,2.10,cw,3.65,th,'panel_alt' if j%2 else 'panel',True)
            label(sl,m.value,x+.21,2.66,cw-.43,.95,th,34 if len(m.value)<12 else 26,True,'accent')
            label(sl,m.label,x+.21,3.77,cw-.43,.67,th,17,True)
            if m.context:label(sl,m.context,x+.21,4.53,cw-.42,.76,th,12,color='secondary')
        insight=spec.key_message or (body[0] if body else '')
        if insight:label(sl,insight,.91,6.36,11.17,.42,th,12,color='accent')
        return
    if mode in ('process','timeline','architecture','methodology','roadmap'):
        items=spec.steps or body
        diagram(sl,items,th,'architecture' if mode=='architecture' else 'timeline' if mode in ('timeline','roadmap') and len(items)<=6 else 'process')
        return
    if mode in ('two_column','problem_solution','comparison','pros_cons'):
        split=max(1,math.ceil(len(body)/2))
        left,right=body[:split],body[split:]
        heads={'problem_solution':('CHALLENGE','RESPONSE'),'pros_cons':('ADVANTAGES','TRADE-OFFS'),'comparison':('OPTION A','OPTION B')}.get(mode,('THE CONTEXT','THE OPPORTUNITY'))
        for idx,col in enumerate((left,right)):
            x=.85+idx*5.88
            rect(sl,x,2.03,5.63,4.25,th,'panel' if idx==0 else 'panel_alt',True)
            label(sl,heads[idx],x+.30,2.37,5.1,.42,th,12,True,'accent')
            wrapped(sl,col,x+.29,3.02,5.06,2.82,th,pt=18)
        return
    if mode in ('three_column','feature_grid','case_study'):
        groups=[body[j:j+2] for j in range(0,len(body),2)] if len(body)>3 else [[s] for s in body]
        groups=groups[:3]
        if not groups:groups=[['Key idea']]
        for j,parts in enumerate(groups):
            cw=11.52/len(groups)-.19;x=.83+j*11.52/len(groups)
            card(sl,th,x,2.13,cw,4.0,j+1,parts[0],parts[1] if len(parts)>1 else '',tint=j%2==1)
        return
    if mode in ('summary','evidence','table','references'):
        if mode=='references':
            wrapped(sl,body,.89,2.09,11.43,3.99,th,pt=17,numbered=True);return
        cells=body[:6]
        for j,text in enumerate(cells):
            cols=2;x=.83+(j%cols)*5.88;y=2.12+(j//cols)*1.38
            rect(sl,x,y,5.64,1.14,th,'panel_alt' if j%2 else 'panel',True)
            label(sl,f'{j+1:02d}',x+.21,y+.15,.48,.24,th,12,True,'accent')
            label(sl,text,x+.72,y+.08,4.65,.93,th,17 if len(text)<105 else 14)
        return
    # bullets, formula and other text-first fallback always remain editable.
    rect(sl,.83,2.05,11.62,4.36,th,'panel',True)
    wrapped(sl,body,.99,2.31,11.18,3.80,th,pt=20)

def build_deck(title:str,slides:list[dict],theme:str='midnight'):
    checked=plan_layouts(validate_deck(slides))
    deck=Presentation();deck.slide_width=Inches(W);deck.slide_height=Inches(H)
    for i,spec in enumerate(checked):
        sl,th=base(deck,i,len(checked),spec,theme)
        render_spec(sl,spec,i,len(checked),th)
        # Preserve supporting content that cannot fit visibly in a visual-led layout.
        # The body and any source references stay in the stored slide JSON too.
        notes=[spec.notes] if spec.notes else []
        if spec.chart or spec.metrics:
            supporting=spec.body if spec.key_message else spec.body[1:]
            if supporting:notes.append('Additional context:\n'+'\n'.join(supporting))
        if spec.steps and spec.body:notes.append('Supporting content:\n'+'\n'.join(spec.body))
        if spec.sources:notes.append('Sources:\n'+'\n'.join(spec.sources))
        if notes:sl.notes_slide.notes_text_frame.text='\n\n'.join(notes)
    return deck,checked

def render(title,slides,theme):
    deck,_=build_deck(title,slides,theme)
    out=BytesIO();deck.save(out);return out.getvalue()
