"""Explainable structural quality gate for editable PowerPoint decks.

Geometric checks are deterministic. Text clipping is an *estimate*, not a
PowerPoint-layout guarantee. Optional LibreOffice PDF conversion verifies that
an independent office renderer can read the exported presentation.
"""
from __future__ import annotations
import io, os, shutil, subprocess, tempfile, math
from pathlib import Path
from pptx import Presentation
from pptx.util import Pt
from pydantic import ValidationError
from pypdf.errors import PyPdfError
from .engine import build_deck


def audit_pptx(data: bytes, *, verify_external: bool=False) -> dict:
    deck=Presentation(io.BytesIO(data))
    defects=[]; slide_results=[]; total_text=0; total_shapes=0; editable=0
    for idx,slide in enumerate(deck.slides,1):
        issues=[]; count_text=0; editable_here=0
        for j,sh in enumerate(slide.shapes):
            total_shapes+=1
            # All generated charts and shapes should remain editable. Pictures are
            # intentionally not counted as native-editable.
            if sh.has_text_frame or sh.has_chart or sh.shape_type==1:
                editable+=1;editable_here+=1
            x,y,w,h=[v/914400 for v in (sh.left,sh.top,sh.width,sh.height)]
            if x<-.015 or y<-.015 or x+w>deck.slide_width/914400+.015 or y+h>deck.slide_height/914400+.015:
                issues.append({'code':'outside_canvas','severity':'critical','element':j,'detail':'Shape extends beyond slide boundaries.'})
            if not sh.has_text_frame or not sh.text.strip():continue
            count_text+=1;total_text+=1
            tf=sh.text_frame
            # Consciously conservative heuristic: emoji/CJK fonts, PowerPoint
            # text engines and corporate font substitutions may differ.
            text=sh.text
            for paragraph in tf.paragraphs:
                if not paragraph.text.strip():continue
                # Most generated labels set their size on the paragraph. Runs
                # inherit it; treating those runs as 18 pt hides small text and
                # measures headings with the wrong dimensions.
                inherited=paragraph.font.size.pt if paragraph.font.size else 18.0
                sizes=[r.font.size.pt if r.font.size else inherited for r in paragraph.runs if r.text]
                fs=max(sizes,default=inherited)
                is_metadata=y<.5 or y>6.86
                smallest=min(sizes,default=inherited)
                if smallest<12 and not is_metadata:
                    issues.append({'code':'small_font','severity':'warning','element':j,'detail':f'{smallest:g} pt may be small for projection.'})
                inner_w=max(.20,w-(tf.margin_left+tf.margin_right)/914400)
                inner_h=max(.12,h-(tf.margin_top+tf.margin_bottom)/914400)
                # Width model errs on warning rather than claiming a definitive clip.
                chars_per_line=max(1,inner_w*72/(fs*.53))
                lines=sum(max(1,math.ceil(len(word)/(chars_per_line))) for word in paragraph.text.split('\n'))
                lines=max(lines,math.ceil(len(paragraph.text)/chars_per_line))
                height_needed=lines*fs*1.22/72
                if height_needed>inner_h*1.34 and not is_metadata:
                    issues.append({'code':'possible_text_overflow','severity':'warning','element':j,'detail':'Text may exceed its box; review rendered output.'})
        if count_text<1:issues.append({'code':'empty_slide','severity':'warning','element':None,'detail':'No text present on this slide.'})
        defects.extend([dict(slide=idx,**item) for item in issues])
        slide_results.append({'slide':idx,'issues':issues,'text_elements':count_text,'editable_elements':editable_here})
    if len(deck.slides)==0:defects.append({'slide':0,'code':'empty_deck','severity':'critical','detail':'Deck contains no slides.'})
    external={'status':'not_requested','pages':None,'detail':None}
    if verify_external:external=verify_render(data,len(deck.slides))
    critical=sum(x['severity']=='critical' for x in defects)
    warning=sum(x['severity']=='warning' for x in defects)
    # Subscores intentionally NOT masquerading as aesthetic or factual validation.
    geometry_score=max(0,100-35*critical-3*warning)
    return {
        'status':'failed' if critical or external['status']=='failed' else 'warnings' if warning else 'structurally_valid',
        'validation_type':'deterministic geometry + heuristic text-size estimates',
        'slide_count':len(deck.slides), 'shape_count':total_shapes,
        'editable_shape_count':editable,'editable_shape_ratio':round(editable/max(total_shapes,1),3),
        'geometry_score_out_of_100':geometry_score, 'score_is_not_visual_quality':True,
        'critical_count':critical,'warning_count':warning,'slides':slide_results,
        'findings':defects[:100], 'independent_render':external,
        'factual_review':'not_performed', 'human_design_review':'not_performed'
    }

def verify_render(data:bytes, expected_pages:int) -> dict:
    executable=shutil.which('libreoffice') or shutil.which('soffice')
    if not executable:
        return {'status':'unavailable','pages':None,'detail':'LibreOffice is not installed in this runtime.'}
    with tempfile.TemporaryDirectory(prefix='pptx-quality-') as temp:
        directory=Path(temp);source=directory/'input.pptx';source.write_bytes(data)
        profile=(directory/'lo-profile').as_uri()
        try:
            completed=subprocess.run([executable,'-env:UserInstallation='+profile,'--headless','--convert-to','pdf','--outdir',temp,str(source)],
                capture_output=True,text=True,timeout=40,check=False,env={**os.environ,'HOME':str(directory)})
            target=directory/'input.pdf'
            if completed.returncode or not target.exists():
                return {'status':'failed','pages':None,'detail':'LibreOffice PDF conversion failed.'}
            from pypdf import PdfReader
            pages=len(PdfReader(str(target)).pages)
            return {'status':'passed' if pages==expected_pages else 'failed','pages':pages,
                'detail':f'Independent conversion produced {pages} of {expected_pages} expected pages.'}
        except (OSError,subprocess.TimeoutExpired,ValueError,PyPdfError) as exc:
            return {'status':'failed','pages':None,'detail':type(exc).__name__}

def audit_slides(title:str,slides:list[dict],theme:str='midnight',verify_external:bool=False)->dict:
    from .engine import render
    return audit_pptx(render(title,slides,theme),verify_external=verify_external)
